"""엑셀 파일로 네트워크(NMS) 장비를 한 번에 여러 대 등록하는 기능.

server.py의 두 라우트(다운로드용 양식 생성 / 업로드된 파일 파싱)에서 이
모듈을 호출한다. 실제 장비 생성은 여기서 하지 않는다 -- 한 대씩 등록할 때와
완전히 같은 검증/연결 테스트를 거치도록, server.py의 _register_device_core()를
행(row) 하나당 한 번씩 그대로 재사용한다 (별도의, 점점 어긋나게 될 두 번째
검증 로직을 만들지 않기 위해서다).
"""
import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TEMPLATE_HEADERS = ['장비명*', 'IP*', '유형', '그룹', '모드 (ping 또는 snmp)', 'SNMP Community', 'SNMP 포트', 'SNMP 버전 (v1 또는 v2c)']
TEMPLATE_EXAMPLE_ROWS = [
    ['3층 스위치', '192.168.0.10', 'L2 Switch', '3층', 'snmp', 'public', 161, 'v2c'],
    ['로비 공유기', '192.168.0.1', '공유기', '', 'ping', '', '', ''],
]
# SNMPv3 자격증명은 필드 수가 많고(사용자명/보안수준/인증·암호화 프로토콜·
# 비밀번호) 엑셀 한 행에 욱여넣기엔 오히려 실수를 유발하기 쉬워서, 일괄
# 등록은 ping과 SNMP v1/v2c까지만 지원한다. v3나 Agent 방식이 필요한 장비는
# 기존처럼 장비 등록 화면에서 하나씩 등록해야 한다 -- 양식 안내문에도 같은
# 내용을 적어둔다.
SUPPORTED_MODES = ('ping', 'snmp')
SUPPORTED_SNMP_VERSIONS = ('v1', 'v2c')


def _style_header_row(ws, header_count):
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4C3DB8', end_color='4C3DB8', fill_type='solid')
    for col_idx in range(1, header_count + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    widths = [18, 16, 14, 10, 20, 16, 12, 22]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def build_template_xlsx():
    """Returns (bytes, filename). Called by GET /download/network-bulk-template."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '네트워크 장비 일괄등록'

    ws.append(TEMPLATE_HEADERS)
    _style_header_row(ws, len(TEMPLATE_HEADERS))
    for row in TEMPLATE_EXAMPLE_ROWS:
        ws.append(row)

    notes = wb.create_sheet('안내')
    notes.append(['네트워크 장비 일괄 등록 안내'])
    notes['A1'].font = Font(bold=True, size=13)
    for i, line in enumerate([
        '',
        '· 장비명, IP는 필수입니다. 나머지는 비워두면 기본값(유형: 기타, 모드: ping, SNMP 포트: 161, SNMP 버전: v2c)이 적용됩니다.',
        '· 모드는 ping 또는 snmp만 지원합니다 (Agent 방식은 이 양식으로 등록할 수 없습니다 -- PC에 프로그램 설치가 필요해서 장비 등록 화면에서 개별 등록해주세요).',
        '· SNMP 버전은 v1 또는 v2c만 지원합니다 (v3는 사용자명/인증/암호화 비밀번호 등 항목이 많아 양식에 담기 어려워 미지원입니다 -- 장비 등록 화면에서 개별 등록해주세요).',
        '· 모드가 snmp인 행은 Community/포트/버전으로 실제 연결 테스트를 거친 뒤에만 등록됩니다 (한 대씩 등록할 때와 동일).',
        '· 모드가 ping인 행은 실제로 핑에 응답하는지 확인한 뒤에만 등록됩니다.',
        '· "네트워크 장비 일괄등록" 시트의 1행은 제목 행이니 지우지 말고, 2행부터 실제 데이터를 입력해주세요 (예시 행은 지우고 입력하셔도 됩니다).',
    ], start=2):
        notes.cell(row=i, column=1, value=line)
    notes.column_dimensions['A'].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), 'InfraSight-네트워크장비-일괄등록양식.xlsx'


def parse_xlsx(file_bytes):
    """Reads an uploaded workbook and returns a list of row dicts, one per
    non-empty data row: {rowNum, name, ip, nettype, group, mode, community,
    snmpPort, snmpVersion, parseError}. rowNum is the 1-based Excel row
    number (so error messages match what the person sees when they open the
    file back up). parseError is set (and every other field left as-is) when
    the row's mode/snmpVersion value isn't one of the supported options --
    everything else is deliberately left for _register_device_core's own
    validation (IP format, name charset, etc.) rather than duplicated here.

    Raises ValueError for anything that means the file itself can't be read
    at all (not an xlsx, empty, wrong sheet) -- the route turns that into a
    400 with the message as-is."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        raise ValueError('엑셀 파일(.xlsx)을 읽을 수 없습니다. InfraSight 양식을 내려받아 그 형식 그대로 입력해주세요.')
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError('빈 파일입니다.')
    # 첫 행은 제목 행으로 간주하고 건너뛴다 -- 정확한 헤더 문구 일치를
    # 요구하지 않는다 (사람이 셀 내용을 살짝 고쳐도 계속 쓸 수 있도록).
    data_rows = rows[1:]
    results = []
    for i, row in enumerate(data_rows, start=2):
        if row is None or all(c is None or str(c).strip() == '' for c in row):
            continue  # 완전히 빈 행은 조용히 건너뛴다 (양식 끝의 여분 행 등)
        cells = list(row) + [None] * (8 - len(row))
        name, ip, nettype, group, mode_raw, community, snmp_port, snmp_version_raw = cells[:8]
        mode = (str(mode_raw).strip().lower() if mode_raw not in (None, '') else 'ping')
        snmp_version = (str(snmp_version_raw).strip().lower() if snmp_version_raw not in (None, '') else 'v2c')
        parse_error = None
        if mode not in SUPPORTED_MODES:
            parse_error = f"모드 값 '{mode_raw}'은(는) 지원하지 않습니다 (ping 또는 snmp만 가능)"
        elif mode == 'snmp' and snmp_version not in SUPPORTED_SNMP_VERSIONS:
            parse_error = f"SNMP 버전 '{snmp_version_raw}'은(는) 지원하지 않습니다 (v1 또는 v2c만 가능, v3는 개별 등록해주세요)"
        results.append({
            'rowNum': i,
            'name': str(name).strip() if name not in (None, '') else '',
            'ip': str(ip).strip() if ip not in (None, '') else '',
            'nettype': str(nettype).strip() if nettype not in (None, '') else '',
            'group': str(group).strip() if group not in (None, '') else '',
            'mode': mode,
            'community': str(community).strip() if community not in (None, '') else '',
            'snmpPort': snmp_port,
            'snmpVersion': snmp_version,
            'parseError': parse_error,
        })
    if not results:
        raise ValueError('입력된 장비 데이터가 없습니다 (2행부터 입력해주세요).')
    return results


def build_export_xlsx(devices):
    """Returns (bytes, filename). devices: list of dicts with
    name/ip/type/group/mode/community/snmpPort/snmpVersion -- the caller
    (server.py) is responsible for pulling the actual device rows + their
    decrypted SNMP community credentials and filtering out anything this
    export can't represent (SNMPv3, Agent mode -- same scope boundary as
    the bulk-import template, see its own SUPPORTED_MODES/
    SUPPORTED_SNMP_VERSIONS comment above).

    Same 8-column layout as build_template_xlsx() on purpose: the file this
    produces can be fed straight into POST /api/devices/bulk-import/network
    on a different InfraSight install with no editing, which is the whole
    point -- "move my currently-registered network devices onto another
    PC's monitoring in one shot"."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '네트워크 장비 내보내기'

    ws.append(TEMPLATE_HEADERS)
    _style_header_row(ws, len(TEMPLATE_HEADERS))
    for d in devices:
        ws.append([
            d['name'], d['ip'], d.get('type') or '', d.get('group') or '',
            d['mode'], d.get('community') or '', d.get('snmpPort') or '', d.get('snmpVersion') or '',
        ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), 'InfraSight-네트워크장비-내보내기.xlsx'
