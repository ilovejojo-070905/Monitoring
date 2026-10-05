"""3-2 보고서 내보내기 -- PDF (reportlab) and Excel (openpyxl) generation for
the 3 report types the ticket asks for: 장애 이력, 자원 사용률, 네트워크 상태.

Korean rendering (PDF): reportlab's built-in fonts are Latin-1 only, so
Correct 한글 output requires registering a real CJK-capable TTF -- 맑은 고딕
(Malgun Gothic), which ships with every Windows install since Windows 7, at
C:\\Windows\\Fonts\\malgun.ttf (+ malgunbd.ttf for bold). This was verified
end-to-end during development: generate a PDF, extract its text back out
with a PDF reader, and confirm the Korean round-trips character-for-character
rather than becoming boxes/mojibake -- not just that the file "looks okay".

Table page-breaks: built via reportlab's platypus layer (SimpleDocTemplate +
Table), not the raw low-level canvas API -- platypus's Table.split() handles
pagination automatically when a table doesn't fit on one page, including
repeating the header row on each continuation page (`repeatRows=1`). Hand-
rolling pagination against the canvas API directly would mean re-implementing
exactly this, with much more room for a subtle bug (a row silently clipped at
a page boundary) that's easy to miss without opening the output in a real PDF
reader. This too was verified: an 80-row test table split cleanly across 3
pages with the header repeated on each.
"""
import io
import os
import time

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import storage
from collector import sla

_FONT_REGISTERED = False
# Only Windows paths -- this whole app is Windows-only already (IS_WINDOWS
# gates elsewhere throughout storage.py/collector/*), so there's no non-
# Windows fallback to maintain.
_KOREAN_FONT_CANDIDATES = [
    (r'C:\Windows\Fonts\malgun.ttf', r'C:\Windows\Fonts\malgunbd.ttf'),
]


def _ensure_korean_font():
    global _FONT_REGISTERED
    if _FONT_REGISTERED:
        return
    for regular, bold in _KOREAN_FONT_CANDIDATES:
        if os.path.exists(regular) and os.path.exists(bold):
            pdfmetrics.registerFont(TTFont('Korean', regular))
            pdfmetrics.registerFont(TTFont('KoreanBold', bold))
            _FONT_REGISTERED = True
            return
    raise RuntimeError('한글 글꼴(맑은 고딕)을 찾을 수 없어 PDF를 생성할 수 없습니다')


def _fmt_ms(ms):
    return time.strftime('%Y-%m-%d %H:%M', time.localtime(ms / 1000)) if ms else '-'


def _fmt_duration(ms):
    if ms is None:
        return '-'
    total_min = round(ms / 60000)
    days, rem = divmod(total_min, 1440)
    hours, mins = divmod(rem, 60)
    if days:
        return f'{days}일 {hours}시간'
    if hours:
        return f'{hours}시간 {mins}분'
    return f'{mins}분'


def _resolve_devices(device_id=None, group=None):
    devices = storage.load_devices()
    if device_id:
        return [d for d in devices if d['id'] == device_id]
    if group:
        return [d for d in devices if (d.get('fields') or {}).get('group') == group]
    return devices


def _filter_desc(device_id=None, group=None):
    if device_id:
        d = storage.load_device(device_id)
        return f"대상 장비: {d['name']}" if d else f"대상 장비: {device_id}"
    if group:
        return f"대상 그룹: {group}"
    return '대상: 전체 장비'


# ------------------------------------------------------------ 데이터 수집 --
def gather_incident_report(start_ms, end_ms, device_id=None, group=None):
    devices = _resolve_devices(device_id, group)
    device_ids = {d['id'] for d in devices}
    conn = storage.get_db()
    rows = conn.execute(
        'SELECT * FROM incidents WHERE ts>=? AND ts<? ORDER BY ts DESC', (start_ms, end_ms)).fetchall()
    conn.close()
    out = []
    for r in rows:
        if r['device_id'] and r['device_id'] not in device_ids:
            continue
        if not r['device_id'] and (device_id or group):
            continue  # a device-less row (e.g. system-level) doesn't match a device/group filter
        out.append({
            'time': _fmt_ms(r['ts']), 'source': r['source'], 'category': r['category'],
            'severity': {'crit': '위험', 'warn': '주의', 'info': '정보'}.get(r['severity'], r['severity']),
            'status': {'open': '미해결', 'ack': '확인됨', 'resolved': '해결됨'}.get(r['status'], r['status']),
            'duration': _fmt_duration((r['resolved_at'] or end_ms) - r['first_occurred_at']) if r['first_occurred_at'] else '-',
            'message': r['message'],
        })
    columns = ['시각', '장비/출처', '분류', '심각도', '상태', '지속시간', '내용']
    table_rows = [[i['time'], i['source'], i['category'], i['severity'], i['status'], i['duration'], i['message']] for i in out]
    return columns, table_rows


def gather_resource_report(start_ms, end_ms, device_id=None, group=None):
    devices = _resolve_devices(device_id, group)
    columns = ['장비명', '평균 CPU %', '최대 CPU %', '평균 메모리 %', '최대 메모리 %', '평균 디스크 %', '최대 디스크 %']
    rows = []
    for d in devices:
        cpu = storage.summarize_metric_range(d['id'], 'cpu', start_ms, end_ms)
        mem = storage.summarize_metric_range(d['id'], 'mem', start_ms, end_ms)
        disk = storage.summarize_metric_range(d['id'], 'disk', start_ms, end_ms)
        if not (cpu or mem or disk):
            continue  # no resource data for this device in range (e.g. a ping-only device) -- omit rather than show all "-"
        def _a(m, key): return f"{m[key]:.1f}" if m else '-'
        rows.append([d['name'], _a(cpu, 'avg'), _a(cpu, 'max'), _a(mem, 'avg'), _a(mem, 'max'), _a(disk, 'avg'), _a(disk, 'max')])
    return columns, rows


def gather_network_report(start_ms, end_ms, device_id=None, group=None):
    """네트워크 상태 보고서: uptime/SLA figures scoped to category='net'
    devices specifically (switches/routers/APs) -- a device/group filter
    further narrows within that, same as the other two report types."""
    devices = [d for d in _resolve_devices(device_id, group) if d['category'] == 'net']
    columns = ['장비명', 'IP', '가동률 %', '장애 횟수', 'MTTR', 'MTBF', '점검 시간']
    rows = []
    for d in devices:
        r = sla.device_uptime_report(d, start_ms, end_ms)
        rows.append([d['name'], d.get('ip') or '-', f"{r['uptimePct']:.3f}", r['failureCount'],
                     _fmt_duration(r['mttrMs']), _fmt_duration(r['mtbfMs']), _fmt_duration(r['maintenanceMs'])])
    return columns, rows


_REPORT_TITLES = {
    'incident': '장애 이력 보고서', 'resource': '자원 사용률 보고서', 'network': '네트워크 상태 보고서',
}
_REPORT_GATHERERS = {
    'incident': gather_incident_report, 'resource': gather_resource_report, 'network': gather_network_report,
}


def build_pdf(report_type, start_ms, end_ms, device_id=None, group=None):
    _ensure_korean_font()
    columns, rows = _REPORT_GATHERERS[report_type](start_ms, end_ms, device_id, group)
    title = _REPORT_TITLES[report_type]
    period_desc = f"기간: {_fmt_ms(start_ms)} ~ {_fmt_ms(end_ms)}"
    filter_desc = _filter_desc(device_id, group)

    title_style = ParagraphStyle('kr_title', fontName='KoreanBold', fontSize=16, leading=20)
    meta_style = ParagraphStyle('kr_meta', fontName='Korean', fontSize=9.5, leading=13, textColor=colors.HexColor('#555555'))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=16 * mm, leftMargin=14 * mm, rightMargin=14 * mm)
    elements = [
        Paragraph(f'InfraSight {title}', title_style), Spacer(1, 3 * mm),
        Paragraph(period_desc, meta_style), Paragraph(filter_desc, meta_style),
        Paragraph(f"생성 시각: {_fmt_ms(int(time.time() * 1000))}", meta_style),
        Spacer(1, 6 * mm),
    ]
    if not rows:
        elements.append(Paragraph('해당 조건에 데이터가 없습니다.', meta_style))
    else:
        # Column widths split the usable page width (A4 minus margins) evenly,
        # except the free-text 내용/장비명 column (if present) gets extra room.
        usable_mm = 210 - 28
        n = len(columns)
        wide_idx = columns.index('내용') if '내용' in columns else (0 if n else None)
        widths = [usable_mm / n] * n
        if wide_idx is not None and n > 1:
            extra = widths[wide_idx] * (n - 1) * 0.5 / n
            for i in range(n):
                widths[i] = widths[i] - extra if i != wide_idx else widths[i] + extra * (n - 1)
        table_data = [columns] + [[Paragraph(str(c), meta_style) for c in row] for row in rows]
        t = Table(table_data, repeatRows=1, colWidths=[w * mm for w in widths])
        t.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), 'Korean'),
            ('FONTNAME', (0, 0), (-1, 0), 'KoreanBold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1c2333')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#dddddd')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f7f8fb')]),
        ]))
        elements.append(t)
    doc.build(elements)
    return buf.getvalue()


def build_excel(report_type, start_ms, end_ms, device_id=None, group=None):
    columns, rows = _REPORT_GATHERERS[report_type](start_ms, end_ms, device_id, group)
    title = _REPORT_TITLES[report_type]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]  # Excel sheet-name length cap
    ws['A1'] = f'InfraSight {title}'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f"기간: {_fmt_ms(start_ms)} ~ {_fmt_ms(end_ms)}"
    ws['A3'] = _filter_desc(device_id, group)
    ws['A4'] = f"생성 시각: {_fmt_ms(int(time.time() * 1000))}"
    header_row = 6
    for ci, col in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=ci, value=col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1C2333')
        cell.alignment = Alignment(horizontal='center')
    for ri, row in enumerate(rows, start=header_row + 1):
        for ci, val in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=val)
    for ci, col in enumerate(columns, start=1):
        max_len = max([len(str(col))] + [len(str(row[ci - 1])) for row in rows]) if rows else len(str(col))
        ws.column_dimensions[get_column_letter(ci)].width = min(max(max_len + 2, 10), 60)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_report(report_type, fmt, start_ms, end_ms, device_id=None, group=None):
    if report_type not in _REPORT_GATHERERS:
        raise ValueError('지원하지 않는 보고서 종류입니다')
    title = _REPORT_TITLES[report_type]
    date_tag = time.strftime('%Y%m%d', time.localtime(start_ms / 1000)) + '-' + time.strftime('%Y%m%d', time.localtime(end_ms / 1000))
    if fmt == 'pdf':
        data = build_pdf(report_type, start_ms, end_ms, device_id, group)
        return data, f'InfraSight_{title}_{date_tag}.pdf', 'application/pdf'
    elif fmt == 'excel':
        data = build_excel(report_type, start_ms, end_ms, device_id, group)
        return data, f'InfraSight_{title}_{date_tag}.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    raise ValueError('지원하지 않는 형식입니다 (pdf 또는 excel)')
